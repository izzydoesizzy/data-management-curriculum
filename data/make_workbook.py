"""
Builds data/powerbi-starter.xlsx: the clean tables (plus the Caseworks-track referrals and appointments), each as a named Excel
Table, ready to upload to the Power BI service (My workspace > New item >
Semantic model > Excel) or open in Power BI Desktop.

Run after generate.py:  python3 data/make_workbook.py   (needs: pip install openpyxl)
"""
import csv
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
SHEETS = [("Services", "services_clean.csv"), ("Clients", "clients.csv"),
          ("ServiceCodes", "service_codes.csv"), ("Workers", "workers.csv"),
          ("Referrals", "referrals.csv"), ("Appointments", "appointments.csv")]
DATE_COLS = {"ServiceDate", "IntakeDate", "DischargeDate", "ReferralDate", "ApptDate"}
NUM_COLS = {"DurationMins": int, "UnitCost": float}

wb = Workbook()
wb.remove(wb.active)
for name, fname in SHEETS:
    ws = wb.create_sheet(name)
    with open(os.path.join(HERE, fname), encoding="utf-8") as f:
        rows = list(csv.reader(f))
    header = rows[0]
    ws.append(header)
    for r in rows[1:]:
        out = []
        for col, v in zip(header, r):
            if v == "":
                out.append(None)
            elif col in DATE_COLS:
                out.append(date.fromisoformat(v))
            elif col in NUM_COLS:
                out.append(NUM_COLS[col](v))
            else:
                out.append(v)
        ws.append(out)
    for i, col in enumerate(header, 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = max(12, len(col) + 4)
        if col in DATE_COLS:
            for cell in ws[letter][1:]:
                cell.number_format = "yyyy-mm-dd"
        if col == "UnitCost":
            for cell in ws[letter][1:]:
                cell.number_format = "$#,##0.00"
    ref = f"A1:{get_column_letter(len(header))}{len(rows)}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = "A2"

wb.properties.creator = "Data Management Curriculum"
wb.properties.created = wb.properties.modified = __import__("datetime").datetime(2026, 1, 1)
wb.save(os.path.join(HERE, "powerbi-starter.xlsx"))
print("wrote powerbi-starter.xlsx")
